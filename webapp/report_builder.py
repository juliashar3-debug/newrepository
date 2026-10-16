"""Rebuild both report formats from an immutable journal snapshot."""
from __future__ import annotations

from copy import deepcopy
from functools import lru_cache
from importlib.util import find_spec
from io import BytesIO
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile


class BuildError(RuntimeError):
    pass


def next_quarter(quarter):
    year, number = quarter.split('-Q')
    return f'{int(year) + (number == "4")}-Q{int(number) % 4 + 1}'


def proposed_plan(quarter):
    target = next_quarter(quarter)
    season = int(target[-1])
    activities = {
        1: ('Доступный зимний досуг', 'Подготовить командную зимнюю прогулку:\nмаршрут и нагрузку выбрать через опрос.'),
        2: ('Знакомство с культурой города', 'Организовать весеннюю экскурсию по Туле:\nтему и дату выбрать вместе с участниками.'),
        3: ('Совместный активный отдых', 'Подготовить летний командный выезд:\nвыбрать доступный маршрут и собрать заявки.'),
        4: ('Обмен результатами работы СМС', 'Провести итоговую встречу участников:\nобсудить опыт и предложения на новый год.'),
    }
    pairs = [
        ('Обмен практическими навыками', 'Провести встречу «Инструменты в работе»:\nкороткие разборы от молодых специалистов.'),
        activities[season],
        ('Совместная помощь по запросу', 'Согласовать с партнером один запрос на помощь\nи собрать добровольцев для общей акции.'),
        ('Участие молодежи в планировании', 'Собрать отзывы после мероприятий\nи уточнить календарь следующего сезона.'),
    ]
    return {'report_quarter': quarter, 'target_quarter': target, 'status': 'proposed',
            'goals_and_tasks': [{'goal': goal, 'task': task.replace('\n', ' ')} for goal, task in pairs]}


def set_period(shape, quarter):
    year, number = quarter.split('-Q')
    roman = ['I','II','III','IV'][int(number)-1]
    text = re.sub(r'\b(?:IV|III|II|I)(?=\s+КВАРТАЛ)', roman, shape.text)
    text = re.sub(r'\d{4}(?=\s+ГОДА)', year, text)
    if text != shape.text:
        set_text(shape, text)


def configure_periods(slides, snapshot):
    quarter = snapshot['quarter']
    for slide in slides[:-1]:
        for shape in slide.shapes:
            if shape.has_text_frame and 'КВАРТАЛ' in shape.text and ('ВЫПОЛНЕН' in shape.text or 'МЕРОПРИЯТИЯ СОВЕТА' in shape.text or re.fullmatch(r'[IV]+ КВАРТАЛ \d{4} ГОДА', shape.text)):
                set_period(shape, quarter)
    final = slides[-1]
    for shape in final.shapes:
        if shape.has_text_frame and 'ЦЕЛИ И ЗАДАЧИ' in shape.text:
            set_period(shape, next_quarter(quarter))
    plan = snapshot['plan']
    for index, pair in enumerate(plan['goals_and_tasks'][:4]):
        for slot, key in [(19 + index * 11,'goal'), (63 + index * 11,'task')]:
            shape = final.shapes[slot]
            text = f'{index+1} {pair[key]}'
            if ' '.join(shape.text.split()) == ' '.join(text.split()):
                continue
            size = shape.text_frame.paragraphs[0].runs[0].font.size
            points = size.pt if size else 8
            lines = wrapped(text, shape.width / 12700 * 12 / points)
            set_text(shape, '\n'.join(lines))
    for index in range(len(plan['goals_and_tasks']),4):
        for slot in [19 + index * 11,63 + index * 11]:
            set_text(final.shapes[slot],'')
    final.notes_slide.notes_text_frame.text = json.dumps(plan, ensure_ascii=False, indent=2)
    if quarter != snapshot['template_quarter']:
        # Photos belong to their reporting period. A new cover stays empty until
        # the owner explicitly supplies that quarter's cover photos.
        cover = slides[0]
        for index, shape in enumerate(list(cover.shapes)[:2]):
            photos = snapshot.get('cover_photos', [])
            if index < len(photos):
                write_picture(cover, shape, Path(photos[index]), 'Фото титульного слайда')
            else:
                shape._element.getparent().remove(shape._element)


def converter():
    candidates = [os.environ.get('SMS_SOFFICE'), shutil.which('soffice'), shutil.which('libreoffice'),
                  '/Applications/LibreOffice.app/Contents/MacOS/soffice']
    for base in ['PROGRAMFILES', 'PROGRAMFILES(X86)']:
        if os.environ.get(base):
            candidates.append(str(Path(os.environ[base]) / 'LibreOffice/program/soffice.exe'))
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return 'libreoffice', str(Path(candidate).resolve())
    if sys.platform == 'win32':
        import winreg
        try:
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, 'PowerPoint.Application/CLSID'.replace('/', '\\')):
                pass
            powershell = shutil.which('powershell.exe')
            if powershell:
                return 'powerpoint', powershell
        except OSError:
            pass
    return None


def capabilities():
    missing = [package for module, package in [('pptx', 'python-pptx'), ('PIL', 'Pillow'), ('pypdf', 'pypdf'), ('fontTools', 'fonttools')]
               if find_spec(module) is None]
    office = converter()
    message = ''
    if missing:
        message = 'Для обновления отчета установите библиотеки: python -m pip install -r webapp/requirements.txt.'
    elif not office:
        message = 'Для создания PDF установите LibreOffice (libreoffice.org). В Windows также поддерживается установленный Microsoft PowerPoint. После установки повторите обновление.'
    return {'ready': not missing and bool(office), 'message': message,
            'converter': office[0] if office else None}


@lru_cache(maxsize=1)
def text_font():
    from PIL import ImageFont
    candidates = [os.environ.get('SMS_REPORT_FONT'), 'C:/Windows/Fonts/verdana.ttf',
                  '/Library/Fonts/Verdana.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    # Measure at 10x slide points to avoid integer rounding on small text.
    for path in candidates:
        if path and Path(path).is_file():
            return ImageFont.truetype(path, 120)
    return ImageFont.load_default(size=120)


def wrapped(text, width=529):
    font = text_font()
    result = []
    for paragraph in text.split('\n'):
        line = ''
        for word in paragraph.split():
            if line and font.getlength(line + ' ' + word) / 10 > width:
                result.append(line)
                line = ''
            while font.getlength(word) / 10 > width:
                split = 1
                while split < len(word) and font.getlength(word[:split+1]) / 10 <= width:
                    split += 1
                result.append(word[:split])
                word = word[split:]
            line = (line + ' ' + word).strip()
        result.append(line)
    return result or ['—']


@lru_cache(maxsize=4)
def font_coverage(bold, italic):
    """Embedded/subset fonts can lack newly entered Cyrillic letters or digits."""
    from fontTools.ttLib import TTFont, TTLibError
    variant = 'z' if bold and italic else 'b' if bold else 'i' if italic else ''
    candidates = [Path('C:/Windows/Fonts') / ('verdana' + variant + '.ttf')]
    if shutil.which('fc-match'):
        style = 'Bold Italic' if bold and italic else 'Bold' if bold else 'Italic' if italic else 'Regular'
        try:
            matched = subprocess.run(['fc-match', '-f', '%{file}', f'Verdana:style={style}'],
                                     capture_output=True, text=True, timeout=5)
            if matched.returncode == 0 and matched.stdout:
                candidates.append(Path(matched.stdout))
        except (OSError, subprocess.TimeoutExpired):
            pass
    for candidate in candidates:
        if candidate.is_file():
            try:
                with TTFont(candidate) as font:
                    return set(font.getBestCmap() or {})
            except (OSError, TTLibError, ImportError):
                pass
    return None


def set_text(shape, text):
    frame = shape.text_frame
    paragraph = frame.paragraphs[0]
    ppr = deepcopy(paragraph._p.pPr)
    rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None
    frame.clear()
    for index, line in enumerate(text.split('\n')):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if para._p.pPr is not None:
            para._p.remove(para._p.pPr)
        if ppr is not None:
            para._p.insert(0, deepcopy(ppr))
        for token in re.findall(r'\s+|\S+', line) or ['']:
            run = para.add_run()
            if rpr is not None:
                run._r.insert(0, deepcopy(rpr))
            run.text = token
            supported = font_coverage(bool(run.font.bold), bool(run.font.italic))
            run.font.name = 'Liberation Sans' if supported is not None and any(ord(char) not in supported for char in token) else 'Verdana'


def clone_slide(presentation, source):
    slide = presentation.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)
    mapping = {}
    for rid, relation in source.part.rels.items():
        if relation.reltype.rsplit('/', 1)[-1] in ['slideLayout', 'notesSlide']:
            continue
        target = relation.target_ref if relation.is_external else relation.target_part
        mapping[rid] = slide.part.relate_to(target, relation.reltype, relation.is_external)
    for shape in source.shapes:
        element = deepcopy(shape._element)
        for node in element.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and value in mapping:
                    node.set(attr, mapping[value])
        slide.shapes._spTree.insert_element_before(element, 'p:extLst')
    return slide


def attendance(event, field):
    value = event.get(field)
    if value is None:
        return '— человек'
    word = 'человека' if not 11 <= value % 100 <= 14 and value % 10 in [2,3,4] else 'человек'
    return f'{value} {word}'


def content(event):
    source = event.get('source', {})
    prior = source.get('report_content', {})
    values = [event['name'], event['description'] or '—', event['organizer'] or '—',
              attendance(event, 'sms_attendance'),
              attendance(event, 'total_attendance'), event['effect'] or '—']
    comparisons = [event['name'] == source.get('name'),
                   event['description'] == source.get('description_draft') and
                   event['date'] == (source.get('date') or '') and event['date_end'] == (source.get('date_end') or ''),
                   event['organizer'] == (source.get('organizer') or ''),
                   event.get('sms_attendance') == source.get('sms_attendance') and not event.get('dirty'),
                   event.get('total_attendance') == (source.get('total_attendance') if source.get('total_attendance') is not None else source.get('visible_people_count')) and not event.get('dirty'),
                   event['effect'] == source.get('effect_draft')]
    for index, key in enumerate(['name', 'description', 'organizer_display', 'sms_attendance_display', 'team_attendance_display', 'effect']):
        if index not in (3,4) and comparisons[index] and key in prior:
            values[index] = prior[key]
    if not comparisons[1]:
        dates = event['date'] or '—'
        if event['date_end'] and event['date_end'] != event['date']:
            dates += ' — ' + event['date_end']
        values[1] += '\nДата проведения: ' + dates + '.'
    return values


def replace_photos(slide, event, files):
    # Preserve the exact approved crop when the event's photo selection is unchanged.
    source = event.get('source', {})
    originals = source.get('photos', [])
    if event['id'] == 'hero_race':
        originals = [{'repository_path': item['repository_path'],
                      'collage_position': 'center' if item['name'].endswith('-3.jpeg') else ''}
                     for item in source.get('photo_repository_manifest', [])]
    original_paths = [p.get('repository_path') for p in originals if p.get('repository_path')]
    actual_paths = [files[p['id']]['repository_path'] for p in event['photos']]
    original_group = next((p.get('repository_path') for p in originals if p.get('collage_position') == 'center'), None)
    if not original_group and original_paths:
        original_group = original_paths[0]
    group = next((p for p in event['photos'] if p.get('group')), None)
    selected_group = files[group['id']]['repository_path'] if group else None
    if event.get('template_matches') and original_paths == actual_paths and original_group == selected_group:
        return
    group = group or (event['photos'][0] if event['photos'] else None)
    others = [p for p in event['photos'] if p is not group]
    slots = [others[0] if others else None, group, others[1] if len(others) > 1 else None]
    for shape, photo in zip(list(slide.shapes)[:3], slots):
        if not photo:
            shape._element.getparent().remove(shape._element)
            continue
        path = Path(files[photo['id']]['path'])
        write_picture(slide, shape, path, photo['name'] + (' — общее фото' if photo is group else ''))


def write_picture(slide, shape, path, name):
    from PIL import Image, ImageOps
    from pptx.oxml.ns import qn
    try:
        with Image.open(path) as original:
            image = ImageOps.exif_transpose(original).convert('RGB')
            image.thumbnail((2000, 2000), Image.Resampling.LANCZOS)
            width, height = image.size
            stream = BytesIO()
            image.save(stream, 'JPEG', quality=94, optimize=True)
    except (OSError, ValueError) as exc:
        raise BuildError('Не удалось прочитать фотографию «' + path.name + '». Загрузите исправный JPG, PNG или WebP.') from exc
    stream.seek(0)
    _, rid = slide.part.get_or_add_image_part(stream)
    shape._element.blipFill.blip.set(qn('r:embed'), rid)
    shape.name = name
    shape._element.nvPicPr.cNvPr.set('descr', name)
    shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
    ratio, image_ratio = shape.width / shape.height, width / height
    if image_ratio > ratio:
        shape.crop_left = shape.crop_right = (1 - ratio / image_ratio) / 2
    else:
        shape.crop_top = shape.crop_bottom = (1 - image_ratio / ratio) / 2


def event_pages(presentation, source_slide, event, files):
    from pptx.util import Pt
    values = content(event)
    indices = [17, 29, 41, 53, 65, 77]
    lines = [wrapped(value) for value in values]
    # Most edits fit the existing six text boxes; retain every original coordinate.
    fits = all(len(value) * 14.4 <= source_slide.shapes[index].height / 12700 + 1.7
               for index, value in zip(indices, lines))
    if all(source_slide.shapes[index].text == value for index, value in zip(indices, values)):
        fits = True
    if fits:
        slide = clone_slide(presentation, source_slide)
        for index, value, parts in zip(indices, values, lines):
            if slide.shapes[index].text != value:
                set_text(slide.shapes[index], '\n'.join(parts))
        replace_photos(slide, event, files)
        pages = [slide]
    else:
        # Long descriptions continue onto additional slides instead of being clipped.
        remaining = [(index, list(parts)) for index, parts in enumerate(lines[1:], 1)]
        pages = []
        while remaining:
            slide = clone_slide(presentation, source_slide)
            rows = [(0, lines[0])]
            available = 327 - (22 + len(lines[0]) * 14.4 + 9)
            while remaining and available >= 45.4:
                index, parts = remaining[0]
                count = min(len(parts), int((available - 31) / 14.4))
                rows.append((index, parts[:count]))
                available -= 31 + count * 14.4
                if count == len(parts):
                    remaining.pop(0)
                else:
                    remaining[0] = (index, parts[count:])
            y = 205.0
            shapes = list(slide.shapes)
            active = {index for index, _ in rows}
            for index, parts in rows:
                group = shapes[6 + index * 12:18 + index * 12]
                label, value = group[-2:]
                delta = Pt(y) - label.top
                for shape in group:
                    shape.top += delta
                value.top = Pt(y + 22)
                value.height = Pt(len(parts) * 14.4 + 1)
                set_text(value, '\n'.join(parts))
                y += 31 + len(parts) * 14.4
            for index in set(range(6)) - active:
                for shape in shapes[6 + index * 12:18 + index * 12]:
                    shape._element.getparent().remove(shape._element)
            replace_photos(slide, event, files)
            pages.append(slide)
    for index, slide in enumerate(pages):
        slide.notes_slide.notes_text_frame.text = json.dumps({
            'event_id': event['id'], 'part': index + 1, 'parts': len(pages),
            'date': event['date'] or None, 'date_end': event['date_end'] or None,
            'formula': event['formula'], 'notes': event['notes'],
            'sms_type': event['sms_type'], 'total_type': event['total_type'],
            'source': event.get('source', {}),
            'effect_status': 'Качественная формулировка; измеримые результаты не предоставлены.'
        }, ensure_ascii=False, indent=2)
    return pages


def export_pdf(pptx, pdf, office):
    pptx, pdf = Path(pptx).resolve(), Path(pdf).resolve()
    program = 'PowerPoint' if office[0] == 'powerpoint' else 'LibreOffice'
    try:
        if office[0] == 'libreoffice':
            with tempfile.TemporaryDirectory(prefix='office-', dir=pptx.parent) as profile:
                result = subprocess.run([office[1], '-env:UserInstallation=' + Path(profile).as_uri(),
                                         '--headless', '--convert-to', 'pdf:impress_pdf_Export',
                                         '--outdir', str(pdf.parent), str(pptx)],
                                        stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180)
        else:
            result = subprocess.run([office[1], '-STA', '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass',
                                     '-File', str(Path(__file__).with_name('export-pdf.ps1')),
                                     str(pptx), str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180)
    except subprocess.TimeoutExpired as exc:
        raise BuildError(f'{program} не завершил создание PDF за три минуты. Проверьте, нет ли в программе окна активации или другого диалога, и повторите обновление. Предыдущие файлы сохранены.') from exc
    except OSError as exc:
        raise BuildError(f'Не удалось запустить {program} для создания PDF: {exc}. Предыдущие файлы сохранены.') from exc
    if result.returncode or not pdf.is_file() or not pdf.read_bytes().startswith(b'%PDF'):
        raise BuildError(export_error(result, office[0]))


def export_error(result, program):
    """Keep COM diagnostics readable on Russian Windows, including PowerShell 5."""
    output = '\n'.join(value.decode('utf-8', errors='replace') if isinstance(value, bytes) else value or ''
                       for value in [result.stdout, result.stderr])
    detail = ''
    for line in output.splitlines():
        if line.startswith('SMS_POWERPOINT_ERROR='):
            try:
                error = json.loads(line.split('=', 1)[1])
                step = {'start': 'запуск', 'open': 'открытие презентации', 'save': 'сохранение PDF'}.get(error.get('stage'), 'экспорт')
                detail = f"{step}, {error.get('hresult', '')}: {error.get('message', '')}"
            except (ValueError, AttributeError):
                continue
            break
    if not detail:
        detail = ' '.join(output.split()) or f'код завершения {result.returncode}; корректный PDF не получен'
    detail = ' '.join(detail.split())[:1200]
    if program == 'powerpoint':
        return (f'PowerPoint не создал PDF. Причина: {detail}. '
                'Откройте PowerPoint вручную, завершите активацию или первый запуск, если программа их предлагает, '
                'и закройте открытые диалоги. Затем повторите сборку. Предыдущие файлы сохранены.')
    return f'LibreOffice не создал PDF. Причина: {detail}. Повторите сборку. Предыдущие файлы сохранены.'


def build(snapshot, directory, stage=lambda message: None):
    ready = capabilities()
    if not ready['ready']:
        raise BuildError(ready['message'])
    from pptx import Presentation
    from pptx.opc.packuri import PackURI
    from pypdf import PdfReader
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    stage('Собираем слайды и фотографии…')
    presentation = Presentation(snapshot['template'])
    originals = list(presentation.slides)
    first_event = originals[1]
    slides = [originals[0]]
    event_slides = {}
    for event in snapshot['events']:
        source = event.get('source', {}).get('report_slide')
        matched = bool(source and 1 < source < len(originals) - 1 and event.get('source',{}).get('reported_quarter') == snapshot['template_quarter'])
        event['template_matches'] = matched
        event_slides[event['id']] = len(slides) + 1
        slides.extend(event_pages(presentation, originals[source - 1] if matched else first_event, event, snapshot['files']))
    slides.extend(originals[-2:])
    configure_periods(slides, snapshot)
    # Removing a photo shape also needs to remove its unused package relation;
    # otherwise a new quarter would still contain old image files internally.
    for slide in slides:
        used={value for node in slide._element.iter() for attr,value in node.attrib.items()
              if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}')}
        for rid,relation in list(slide.part.rels.items()):
            if relation.reltype.endswith('/image') and rid not in used:
                slide.part.drop_rel(rid)
    # Keep source package parts for the template's logos, curves and editable artwork.
    identifiers = list(presentation.slides._sldIdLst)
    lookup = {presentation.part.related_slide(item.rId).part.partname: item for item in identifiers}
    retained = {slide.part.partname for slide in slides}
    for item in identifiers:
        presentation.slides._sldIdLst.remove(item)
        if presentation.part.related_slide(item.rId).part.partname not in retained:
            presentation.part.drop_rel(item.rId)
    for slide in slides:
        presentation.slides._sldIdLst.append(lookup[slide.part.partname])
    # python-pptx allocates future slide names from the slide count. Keep names
    # contiguous so another update can safely use this report as its template.
    for index, slide in enumerate(slides, 1):
        slide.part._partname = PackURI(f'/ppt/slides/slide{index}.xml')
    year, number = snapshot['quarter'].split('-Q')
    roman = ['I','II','III','IV'][int(number)-1]
    presentation.core_properties.title = f'Отчет СМС АО ЦКБА за {roman} квартал {year} года'
    presentation.core_properties.subject = 'Обновлено по сохраненным карточкам сайта; неизвестные данные оставлены прочерками.'
    pptx = directory / 'report.pptx'
    pdf = directory / 'report.pdf'
    presentation.save(pptx)
    if len(Presentation(pptx).slides) != len(slides):
        raise BuildError('Не удалось проверить состав презентации. Предыдущие файлы сохранены.')
    stage('Создаем PDF из новой презентации…')
    export_pdf(pptx, pdf, converter())
    if len(PdfReader(pdf).pages) != len(slides):
        raise BuildError('Количество страниц PDF не совпало с презентацией. Предыдущие файлы сохранены.')
    stage('Сохраняем новую версию…')
    return {'slide_count': len(slides), 'event_count': len(snapshot['events']), 'event_slides': event_slides}
